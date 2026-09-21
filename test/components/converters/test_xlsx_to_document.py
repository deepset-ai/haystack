# SPDX-FileCopyrightText: 2022-present deepset GmbH <info@deepset.ai>
#
# SPDX-License-Identifier: Apache-2.0

import csv
import io
import logging
from pathlib import Path
from typing import Literal

import pytest
from openpyxl import Workbook

from haystack.components.converters.xlsx import XLSXToDocument
from haystack.dataclasses import ByteStream


class TestXLSXToDocument:
    def test_init(self) -> None:
        converter = XLSXToDocument()
        assert converter.sheet_name is None
        assert converter.read_excel_kwargs == {}
        assert converter.table_format == "csv"
        assert converter.link_format == "none"
        assert converter.table_format_kwargs == {}

    def test_run_basic_tables(self, test_files_path: Path) -> None:
        converter = XLSXToDocument(store_full_path=True)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"]
        results = converter.run(sources=paths, meta={"date_added": "2022-01-01T00:00:00"})
        documents = results["documents"]
        assert len(documents) == 2
        assert documents[0].content == ",A,B\n1,col_a,col_b\n2,1.5,test\n"
        assert documents[0].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": str(test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"),
            "xlsx": {"sheet_name": "Basic Table"},
        }
        assert documents[1].content == ",A,B\n1,col_c,col_d\n2,True,\n"
        assert documents[1].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": str(test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"),
            "xlsx": {"sheet_name": "Table Missing Value"},
        }

    def test_run_table_empty_rows_and_columns(self, test_files_path: Path) -> None:
        converter = XLSXToDocument(store_full_path=False)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "table_empty_rows_and_columns.xlsx"]
        results = converter.run(sources=paths, meta={"date_added": "2022-01-01T00:00:00"})
        documents = results["documents"]
        assert len(documents) == 1
        assert documents[0].content == ",A,B,C\n1,,,\n2,,,\n3,,,\n4,,col_a,col_b\n5,,1.5,test\n"
        assert documents[0].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": "table_empty_rows_and_columns.xlsx",
            "xlsx": {"sheet_name": "Sheet1"},
        }

    def test_run_multiple_tables_in_one_sheet(self, test_files_path: Path) -> None:
        converter = XLSXToDocument(store_full_path=True)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "multiple_tables.xlsx"]
        results = converter.run(sources=paths, meta={"date_added": "2022-01-01T00:00:00"})
        documents = results["documents"]
        assert len(documents) == 1
        assert (
            documents[0].content
            == ",A,B,C,D,E,F\n1,,,,,,\n2,,,,,,\n3,,col_a,col_b,,,\n4,,1.5,test,,col_c,col_d\n5,,,,,3,True\n"
        )
        assert documents[0].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": str(test_files_path / "xlsx" / "multiple_tables.xlsx"),
            "xlsx": {"sheet_name": "Sheet1"},
        }

    def test_run_markdown(self, test_files_path: Path) -> None:
        converter = XLSXToDocument(table_format="markdown", store_full_path=True)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"]
        results = converter.run(sources=paths, meta={"date_added": "2022-01-01T00:00:00"})
        documents = results["documents"]
        assert len(documents) == 2
        assert (
            documents[0].content
            == "|    | A     | B     |\n|---:|:------|:------|\n|  1 | col_a | col_b |\n|  2 | 1.5   | test  |"
        )
        assert documents[0].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": str(test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"),
            "xlsx": {"sheet_name": "Basic Table"},
        }
        assert (
            documents[1].content
            # The empty cell reads as empty, the way the CSV format already writes it.
            == "|    | A     | B     |\n|---:|:------|:------|\n|  1 | col_c | col_d |\n|  2 | True  |       |"
        )
        assert documents[1].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": str(test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"),
            "xlsx": {"sheet_name": "Table Missing Value"},
        }

    def test_run_markdown_missing_value(self, test_files_path: Path) -> None:
        """table_format_kwargs["missingval"] reaches tabulate for an empty cell."""
        converter = XLSXToDocument(table_format="markdown", table_format_kwargs={"missingval": "N/A"})
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"]
        results = converter.run(sources=paths)
        assert (
            results["documents"][1].content
            == "|    | A     | B     |\n|---:|:------|:------|\n|  1 | col_c | col_d |\n|  2 | True  | N/A   |"
        )

    @pytest.mark.parametrize(
        "sheet_name, expected_sheet_name, expected_content",
        [
            ("Basic Table", "Basic Table", ",A,B\n1,col_a,col_b\n2,1.5,test\n"),
            ("Table Missing Value", "Table Missing Value", ",A,B\n1,col_c,col_d\n2,True,\n"),
            (0, 0, ",A,B\n1,col_a,col_b\n2,1.5,test\n"),
            (1, 1, ",A,B\n1,col_c,col_d\n2,True,\n"),
        ],
    )
    def test_run_sheet_name(
        self, sheet_name: int | str, expected_sheet_name: str, expected_content: str, test_files_path: Path
    ) -> None:
        converter = XLSXToDocument(sheet_name=sheet_name, store_full_path=True)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"]
        results = converter.run(sources=paths)
        documents = results["documents"]
        assert len(documents) == 1
        assert documents[0].content == expected_content
        assert documents[0].meta == {
            "file_path": str(test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"),
            "xlsx": {"sheet_name": expected_sheet_name},
        }

    def test_run_with_read_excel_kwargs(self, test_files_path: Path) -> None:
        converter = XLSXToDocument(sheet_name="Basic Table", read_excel_kwargs={"skiprows": 1}, store_full_path=True)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"]
        results = converter.run(sources=paths, meta={"date_added": "2022-01-01T00:00:00"})
        documents = results["documents"]
        assert len(documents) == 1
        assert documents[0].content == ",A,B\n1,1.5,test\n"
        assert documents[0].meta == {
            "date_added": "2022-01-01T00:00:00",
            "file_path": str(test_files_path / "xlsx" / "basic_tables_two_sheets.xlsx"),
            "xlsx": {"sheet_name": "Basic Table"},
        }

    def test_run_error_wrong_file_type(self, caplog: pytest.LogCaptureFixture, test_files_path: Path) -> None:
        converter = XLSXToDocument()
        sources: list[str | Path | ByteStream] = [test_files_path / "pdf" / "sample_pdf_1.pdf"]
        with caplog.at_level(logging.WARNING):
            results = converter.run(sources=sources)
            assert "sample_pdf_1.pdf and convert it" in caplog.text
            assert results["documents"] == []

    def test_run_error_non_existent_file(self, caplog: pytest.LogCaptureFixture) -> None:
        converter = XLSXToDocument()
        paths: list[str | Path | ByteStream] = ["non_existing_file.docx"]
        with caplog.at_level(logging.WARNING):
            converter.run(sources=paths)
            assert "Could not read non_existing_file.docx" in caplog.text

    def test_link_format_invalid(self) -> None:
        with pytest.raises(ValueError, match="Unknown link format"):
            XLSXToDocument(link_format="invalid")  # type: ignore[arg-type]

    @pytest.mark.parametrize("link_format", ["markdown", "plain"])
    def test_link_extraction(self, test_files_path: Path, link_format: Literal["markdown", "plain"]) -> None:
        converter = XLSXToDocument(link_format=link_format)
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "spreadsheet_with_links.xlsx"]
        results = converter.run(sources=paths)
        content = results["documents"][0].content
        assert content is not None

        if link_format == "markdown":
            assert "[Click here](https://example.com)" in content
            assert "[Docs](https://python.org)" in content
        else:
            assert "Click here (https://example.com)" in content
            assert "Docs (https://python.org)" in content

    def test_no_link_extraction(self, test_files_path: Path) -> None:
        converter = XLSXToDocument()
        paths: list[str | Path | ByteStream] = [test_files_path / "xlsx" / "spreadsheet_with_links.xlsx"]
        results = converter.run(sources=paths)
        content = results["documents"][0].content
        assert content is not None

        assert "https://example.com" not in content
        assert "Click here" in content

    @pytest.mark.parametrize("link_format", ["markdown", "plain"])
    @pytest.mark.parametrize("table_format", ["csv", "markdown"])
    @pytest.mark.parametrize("use_bytestream", [False, True])
    def test_link_extraction_from_typed_columns(
        self,
        tmp_path: Path,
        link_format: Literal["markdown", "plain"],
        table_format: Literal["csv", "markdown"],
        use_bytestream: bool,
    ) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append([42, 3.5, True, 1.23456])
        sheet.append([43, 4.5, False, 2.34567])
        for cell in ["A1", "B1", "C1"]:
            sheet[cell].hyperlink = "https://example.com"
        path = tmp_path / "typed_links.xlsx"
        workbook.save(path)
        workbook.close()

        source = ByteStream.from_file_path(path) if use_bytestream else path
        format_kwargs = {"float_format": "%.2f"} if table_format == "csv" else {"floatfmt": ".2f"}
        converter = XLSXToDocument(
            link_format=link_format, table_format=table_format, table_format_kwargs=format_kwargs
        )
        documents = converter.run(sources=[source])["documents"]

        assert len(documents) == 1
        content = documents[0].content
        assert content is not None
        if table_format == "csv":
            rows = [row[1:] for row in list(csv.reader(io.StringIO(content)))[1:]]
        else:
            rows = [[cell.strip() for cell in row.split("|")[2:-1]] for row in content.splitlines()[2:]]
        linked_values = [
            f"[{text}](https://example.com)" if link_format == "markdown" else f"{text} (https://example.com)"
            for text in ["42", "3.5", "True"]
        ]
        assert rows == [linked_values + ["1.23"], ["43", "4.5", "False", "2.35"]]
        assert documents[0].meta["xlsx"] == {"sheet_name": "Sheet"}
