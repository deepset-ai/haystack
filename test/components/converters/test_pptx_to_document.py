# SPDX-FileCopyrightText: 2022-present deepset GmbH <info@deepset.ai>
#
# SPDX-License-Identifier: Apache-2.0

import io
import logging
import os
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches

from haystack.components.converters.pptx import PPTXToDocument
from haystack.dataclasses import ByteStream


class TestPPTXToDocument:
    def test_run(self, test_files_path):
        """
        Test if the component runs correctly.
        """
        bytestream = ByteStream.from_file_path(test_files_path / "pptx" / "sample_pptx.pptx")
        bytestream.meta["file_path"] = str(test_files_path / "pptx" / "sample_pptx.pptx")
        bytestream.meta["key"] = "value"
        first_path = str(test_files_path / "pptx" / "sample_pptx.pptx")
        files: list[str | Path | ByteStream] = [first_path, bytestream]
        converter = PPTXToDocument()
        output = converter.run(sources=files)
        docs = output["documents"]

        assert len(docs) == 2
        assert (
            "Sample Title Slide\nJane Doe\fTitle of First Slide\nThis is a bullet point\nThis is another bullet point"
            in docs[0].content
        )
        assert (
            "Sample Title Slide\nJane Doe\fTitle of First Slide\nThis is a bullet point\nThis is another bullet point"
            in docs[0].content
        )
        assert docs[0].meta["file_path"] == os.path.basename(first_path)
        assert docs[1].meta == {"file_path": os.path.basename(bytestream.meta["file_path"]), "key": "value"}

    def test_run_error_non_existent_file(self, caplog):
        sources: list[str | Path | ByteStream] = ["non_existing_file.pptx"]
        converter = PPTXToDocument()
        with caplog.at_level(logging.WARNING):
            results = converter.run(sources=sources)
            assert "Could not read non_existing_file.pptx" in caplog.text
            assert results["documents"] == []

    def test_run_error_wrong_file_type(self, caplog, test_files_path):
        sources: list[str | Path | ByteStream] = [str(test_files_path / "txt" / "doc_1.txt")]
        converter = PPTXToDocument()
        with caplog.at_level(logging.WARNING):
            results = converter.run(sources=sources)
            assert "doc_1.txt and convert it" in caplog.text
            assert results["documents"] == []

    def test_run_with_meta(self, test_files_path):
        bytestream = ByteStream.from_file_path(test_files_path / "pptx" / "sample_pptx.pptx")
        bytestream.meta["file_path"] = str(test_files_path / "pptx" / "sample_pptx.pptx")
        bytestream.meta["key"] = "value"

        converter = PPTXToDocument()
        output = converter.run(sources=[bytestream], meta=[{"language": "it"}])
        document = output["documents"][0]

        assert document.meta == {
            "file_path": os.path.basename(test_files_path / "pptx" / "sample_pptx.pptx"),
            "key": "value",
            "language": "it",
        }

    def test_run_with_store_full_path_false(self, test_files_path):
        """
        Test if the component runs correctly with store_full_path=False
        """
        bytestream = ByteStream.from_file_path(test_files_path / "pptx" / "sample_pptx.pptx")
        bytestream.meta["file_path"] = str(test_files_path / "pptx" / "sample_pptx.pptx")
        bytestream.meta["key"] = "value"

        converter = PPTXToDocument(store_full_path=False)
        output = converter.run(sources=[bytestream], meta=[{"language": "it"}])
        document = output["documents"][0]

        assert document.meta == {"file_path": "sample_pptx.pptx", "key": "value", "language": "it"}

    def test_to_dict(self):
        converter = PPTXToDocument(link_format="markdown", store_full_path=True)
        data = converter.to_dict()
        assert data == {
            "type": "haystack.components.converters.pptx.PPTXToDocument",
            "init_parameters": {"link_format": "markdown", "store_full_path": True},
        }

    def test_to_dict_defaults(self):
        converter = PPTXToDocument()
        data = converter.to_dict()
        assert data == {
            "type": "haystack.components.converters.pptx.PPTXToDocument",
            "init_parameters": {"link_format": "none", "store_full_path": False},
        }

    def test_link_format_invalid(self):
        with pytest.raises(ValueError, match="Unknown link format"):
            PPTXToDocument(link_format="invalid")  # type: ignore[arg-type]

    @pytest.mark.parametrize("link_format", ["markdown", "plain"])
    def test_link_extraction(self, test_files_path, link_format):
        converter = PPTXToDocument(link_format=link_format)
        paths = [test_files_path / "pptx" / "sample_pptx_with_link.pptx"]
        output = converter.run(sources=paths)
        content = output["documents"][0].content

        if link_format == "markdown":
            assert "[Example](https://example.com)" in content
        else:
            assert "Example (https://example.com)" in content

    def test_no_link_extraction(self, test_files_path):
        converter = PPTXToDocument()
        paths = [test_files_path / "pptx" / "sample_pptx_with_link.pptx"]
        output = converter.run(sources=paths)
        content = output["documents"][0].content

        assert "https://example.com" not in content
        assert "Example" in content

    @pytest.mark.parametrize(
        "link_format,expected",
        [
            ("none", "Slide 3 of the deck"),
            ("markdown", "Slide 3 of [the deck](https://example.com)"),
            ("plain", "Slide 3 of the deck (https://example.com)"),
        ],
    )
    def test_run_keeps_field_text_with_every_link_format(self, link_format, expected):
        # A slide-number field (`a:fld`) sits between the runs of the paragraph
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.paragraphs[0]
        paragraph.add_run().text = "Slide "
        paragraph._p.append(
            parse_xml(
                f'<a:fld {nsdecls("a")} id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
                "<a:t>3</a:t></a:fld>"
            )
        )
        paragraph.add_run().text = " of "
        link = paragraph.add_run()
        link.text = "the deck"
        link.hyperlink.address = "https://example.com"
        buffer = io.BytesIO()
        presentation.save(buffer)

        documents = PPTXToDocument(link_format=link_format).run(sources=[ByteStream(buffer.getvalue())])["documents"]

        assert documents[0].content == expected

    def test_run_reads_nested_groups_and_table_with_links(self, test_files_path):
        # The slide includes nested text groups, a grouped table, and an empty decorative shape.
        converter = PPTXToDocument(link_format="markdown")
        source = test_files_path / "pptx" / "complex_groups_and_table.pptx"

        documents = converter.run(sources=[source])["documents"]

        assert len(documents) == 1
        assert documents[0].content == (
            "PPTX extraction playground\n"
            "One slide • text boxes • nested groups • table cells • hyperlinks\n"
            "PLAIN TEXTBOX\n"
            "First paragraph.\n"
            "Second paragraph: café, naïve, 日本語.\n"
            "[Haystack documentation](https://docs.haystack.deepset.ai/)\n"
            "OUTER GROUP\n"
            "Text directly inside the outer group.\n"
            "NESTED GROUP\n"
            "Text two levels deep.\n"
            "[deepset website](https://www.deepset.ai/)\n"
            "TABLE INSIDE A GROUP\n"
            "Case,Content,Reference\r\n"
            'CSV escaping,"Hello, world\nHe said ""yes""",\r\n'
            "Markdown,left | right,[Docs](https://docs.haystack.deepset.ai/)\r\n"
            'Unicode,café / 日本語,"Two\nparagraphs"\n'
            "\n"
            "END OF SLIDE — text after the groups and table"
        )
