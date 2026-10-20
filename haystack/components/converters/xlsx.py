# SPDX-FileCopyrightText: 2022-present deepset GmbH <info@deepset.ai>
#
# SPDX-License-Identifier: Apache-2.0

import io
import os
from pathlib import Path
from typing import Any, Literal

from haystack import Document, component, logging
from haystack.components.converters.utils import get_bytestream_from_source, normalize_metadata
from haystack.dataclasses import ByteStream
from haystack.lazy_imports import LazyImport

logger = logging.getLogger(__name__)

with LazyImport("Run 'pip install pandas openpyxl'") as pandas_xlsx_import:
    import openpyxl
    import pandas as pd

with LazyImport("Run 'pip install tabulate'") as tabulate_import:
    from tabulate import tabulate  # noqa: F401 # the library is used but not directly referenced

# The `pandas.read_excel` options that decide which cells end up where in the DataFrame.
_POSITION_KWARGS = ("skiprows", "nrows", "skipfooter", "usecols", "names", "index_col")
# Placeholder written into non-empty cells so pandas itself can resolve where they land after the
# positional read_excel_kwargs (skiprows/names/usecols/...) have been applied.
_CELL_TAG = "__haystack_cell__"


@component
class XLSXToDocument:
    """
    Converts XLSX (Excel) files into Documents.

    Supports reading data from specific sheets or all sheets in the Excel file. If all sheets are read, a Document is
    created for each sheet. The content of the Document is the table which can be saved in CSV or Markdown format.

    ### Usage example

    ```python
    from haystack.components.converters.xlsx import XLSXToDocument
    from datetime import datetime

    converter = XLSXToDocument()
    results = converter.run(
        sources=["test/test_files/xlsx/basic_tables_two_sheets.xlsx"], meta={"date_added": datetime.now().isoformat()}
    )
    documents = results["documents"]

    print(documents[0].content)
    # >> ",A,B\\n1,col_a,col_b\\n2,1.5,test\\n"
    ```
    """

    def __init__(
        self,
        table_format: Literal["csv", "markdown"] = "csv",
        sheet_name: str | int | list[str | int] | None = None,
        read_excel_kwargs: dict[str, Any] | None = None,
        table_format_kwargs: dict[str, Any] | None = None,
        *,
        link_format: Literal["markdown", "plain", "none"] = "none",
        store_full_path: bool = False,
    ) -> None:
        """
        Creates a XLSXToDocument component.

        :param table_format: The format to convert the Excel file to.
        :param sheet_name: The name of the sheet to read. If None, all sheets are read.
        :param read_excel_kwargs: Additional arguments to pass to `pandas.read_excel`.
            See https://pandas.pydata.org/docs/reference/api/pandas.read_excel.html#pandas-read-excel
        :param table_format_kwargs: Additional keyword arguments to pass to the table format function.
            - If `table_format` is "csv", these arguments are passed to `pandas.DataFrame.to_csv`.
              See https://pandas.pydata.org/docs/reference/api/pandas.DataFrame.to_csv.html#pandas-dataframe-to-csv
            - If `table_format` is "markdown", these arguments are passed to `pandas.DataFrame.to_markdown`.
              See https://pandas.pydata.org/docs/reference/api/pandas.DataFrame.to_markdown.html#pandas-dataframe-to-markdown
        :param link_format: The format for link output. Possible options:
            - `"markdown"`: `[text](url)`
            - `"plain"`: `text (url)`
            - `"none"`: Only the text is extracted, link addresses are ignored.
        :param store_full_path:
            If True, the full path of the file is stored in the metadata of the document.
            If False, only the file name is stored.
        """
        pandas_xlsx_import.check()
        self.table_format = table_format
        if table_format not in ["csv", "markdown"]:
            raise ValueError(f"Unsupported export format: {table_format}. Choose either 'csv' or 'markdown'.")
        if link_format not in ("markdown", "plain", "none"):
            msg = f"Unknown link format '{link_format}'. Supported formats are: 'markdown', 'plain', 'none'"
            raise ValueError(msg)
        if table_format == "markdown":
            tabulate_import.check()
        self.link_format = link_format
        self.sheet_name = sheet_name
        self.read_excel_kwargs = read_excel_kwargs or {}
        self.table_format_kwargs = table_format_kwargs or {}
        self.store_full_path = store_full_path

    @component.output_types(documents=list[Document])
    def run(
        self, sources: list[str | Path | ByteStream], meta: dict[str, Any] | list[dict[str, Any]] | None = None
    ) -> dict[str, list[Document]]:
        """
        Converts a XLSX file to a Document.

        :param sources:
            List of file paths or ByteStream objects.
        :param meta:
            Optional metadata to attach to the documents.
            This value can be either a list of dictionaries or a single dictionary.
            If it's a single dictionary, its content is added to the metadata of all produced documents.
            If it's a list, the length of the list must match the number of sources, because the two lists will
            be zipped.
            If `sources` contains ByteStream objects, their `meta` will be added to the output documents.
        :returns:
            A dictionary with the following keys:
            - `documents`: Created documents
        """
        documents = []

        meta_list = normalize_metadata(meta, sources_count=len(sources))

        for source, metadata in zip(sources, meta_list, strict=True):
            try:
                bytestream = get_bytestream_from_source(source)
            except Exception as e:
                logger.warning("Could not read {source}. Skipping it. Error: {error}", source=source, error=e)
                continue

            try:
                tables, tables_metadata = self._extract_tables(bytestream)
            except Exception as e:
                logger.warning(
                    "Could not read {source} and convert it to a Document, skipping. Error: {error}",
                    source=source,
                    error=e,
                )
                continue

            # Loop over tables and create a Document for each table
            for table, excel_metadata in zip(tables, tables_metadata, strict=True):
                merged_metadata = {**bytestream.meta, **metadata, **excel_metadata}

                if not self.store_full_path and "file_path" in bytestream.meta:
                    file_path = bytestream.meta["file_path"]
                    merged_metadata["file_path"] = os.path.basename(file_path)

                document = Document(content=table, meta=merged_metadata)
                documents.append(document)

        return {"documents": documents}

    @staticmethod
    def _generate_excel_column_names(n_cols: int) -> list[str]:
        result = []
        for i in range(n_cols):
            col_name = ""
            num = i
            while num >= 0:
                col_name = chr(num % 26 + 65) + col_name
                num = num // 26 - 1
            result.append(col_name)
        return result

    @staticmethod
    def _get_worksheet(workbook: Any, sheet_key: str | int | None) -> Any:
        if isinstance(sheet_key, int):
            return workbook.worksheets[sheet_key]
        if sheet_key is None:
            return workbook.active
        return workbook[sheet_key]

    def _map_hyperlinks_to_dataframe(
        self, workbook: Any, hyperlinks_by_sheet: dict[str | int | None, dict[tuple[int, int], str]]
    ) -> dict[str | int | None, dict[tuple[int, int], str]]:
        """Re-key 0-based Excel cell coordinates to their positions in the filtered DataFrame."""
        for sheet_key in hyperlinks_by_sheet:
            for row in self._get_worksheet(workbook, sheet_key).iter_rows():
                for cell in row:
                    # Empty cells must stay empty: pandas sizes the sheet by its non-empty cells, and the placeholder
                    # cells of merged ranges are read-only.
                    if cell.value not in (None, ""):
                        cell.value = f"{_CELL_TAG}{cell.row - 1},{cell.column - 1}"
        tagged_bytes = io.BytesIO()
        workbook.save(tagged_bytes)
        position_kwargs = {
            key: self.read_excel_kwargs[key] for key in _POSITION_KWARGS if key in self.read_excel_kwargs
        }
        tagged_sheets = pd.read_excel(
            io=tagged_bytes, sheet_name=self.sheet_name, header=None, engine="openpyxl", **position_kwargs
        )
        if isinstance(tagged_sheets, pd.DataFrame):
            tagged_sheets = {self.sheet_name: tagged_sheets}

        mapped_hyperlinks = {}
        for sheet_key, cell_links in hyperlinks_by_sheet.items():
            positions = {}
            for row_idx, values in enumerate(tagged_sheets[sheet_key].itertuples(index=False, name=None)):
                for col_idx, value in enumerate(values):
                    if isinstance(value, str) and value.startswith(_CELL_TAG):
                        excel_row, excel_col = value.removeprefix(_CELL_TAG).split(",")
                        positions[(int(excel_row), int(excel_col))] = (row_idx, col_idx)
            mapped_hyperlinks[sheet_key] = {
                positions[cell]: url for cell, url in cell_links.items() if cell in positions
            }
        return mapped_hyperlinks

    def _extract_tables(self, bytestream: ByteStream) -> tuple[list[str], list[dict]]:
        """
        Extract tables from an Excel file.
        """
        file_bytes = io.BytesIO(bytestream.data)
        resolved_read_excel_kwargs = {
            **self.read_excel_kwargs,
            "sheet_name": self.sheet_name,
            "header": None,  # Don't assign any pandas column labels
            "engine": "openpyxl",  # Use openpyxl as the engine to read the Excel file
        }
        sheet_to_dataframe = pd.read_excel(io=file_bytes, **resolved_read_excel_kwargs)
        if isinstance(sheet_to_dataframe, pd.DataFrame):
            sheet_to_dataframe = {self.sheet_name: sheet_to_dataframe}

        # If link extraction is enabled, load the workbook with openpyxl to read hyperlinks
        hyperlinks_by_sheet: dict[str | int | None, dict[tuple[int, int], str]] = {}
        if self.link_format != "none":
            file_bytes.seek(0)
            wb = openpyxl.load_workbook(file_bytes, data_only=True)
            for sheet_key in sheet_to_dataframe:
                ws = self._get_worksheet(wb, sheet_key)
                cell_links: dict[tuple[int, int], str] = {}
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.hyperlink and cell.hyperlink.target:
                            # Store zero-based Excel coordinates; they are translated to filtered positions below.
                            cell_links[(cell.row - 1, cell.column - 1)] = cell.hyperlink.target
                hyperlinks_by_sheet[sheet_key] = cell_links
            if any(hyperlinks_by_sheet.values()) and any(key in self.read_excel_kwargs for key in _POSITION_KWARGS):
                hyperlinks_by_sheet = self._map_hyperlinks_to_dataframe(wb, hyperlinks_by_sheet)
            wb.close()

        updated_sheet_to_dataframe = {}
        for key in sheet_to_dataframe:
            df = sheet_to_dataframe[key]
            # Row starts at 1 in Excel
            df.index = df.index + 1
            # Excel column names are Alphabet Characters
            header = self._generate_excel_column_names(df.shape[1])
            df.columns = header

            # Apply hyperlinks to cell values
            if key in hyperlinks_by_sheet:
                for (row_idx, col_idx), url in hyperlinks_by_sheet[key].items():
                    if row_idx < len(df) and col_idx < len(df.columns):
                        cell_value = df.iat[row_idx, col_idx]
                        text = str(cell_value) if pd.notna(cell_value) else ""
                        # Hyperlink text must be assignable to numeric and other typed columns.
                        column = df.columns[col_idx]
                        if df[column].dtype != object:
                            df[column] = df[column].astype(object)
                        if self.link_format == "markdown":
                            df.iat[row_idx, col_idx] = f"[{text}]({url})"
                        else:
                            df.iat[row_idx, col_idx] = f"{text} ({url})"

            updated_sheet_to_dataframe[key] = df

        tables = []
        metadata = []
        for key, value in updated_sheet_to_dataframe.items():
            if self.table_format == "csv":
                resolved_kwargs = {"index": True, "header": True, "lineterminator": "\n", **self.table_format_kwargs}
                tables.append(value.to_csv(**resolved_kwargs))
            else:
                resolved_kwargs = {
                    "index": True,
                    "headers": value.columns,
                    "tablefmt": "pipe",
                    "missingval": "",
                    **self.table_format_kwargs,
                }
                if resolved_kwargs["tablefmt"] == "pipe":
                    value = value.replace(
                        {
                            r"\r\n|\r|\n": " ",  # keep in-cell line breaks from creating extra Markdown rows
                            r"(\\*)\|": r"\1\1\\|",  # escape pipes but preserve any preceding literal backslashes
                        },
                        regex=True,
                    )

                # to_markdown uses tabulate, whose missingval only covers None: a NaN
                # reaches the formatter as a number and is written out as "nan". Replace
                # the empty cells with None so an empty cell reads as empty, the way
                # to_csv already writes it, and so missingval keeps working.
                filled = value.astype(object).where(value.notna(), None)
                tables.append(filled.to_markdown(**resolved_kwargs))
            # add sheet_name to metadata
            metadata.append({"xlsx": {"sheet_name": key}})
        return tables, metadata
